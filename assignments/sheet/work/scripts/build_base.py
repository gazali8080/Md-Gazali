"""Base workbook: question text, original data, updated data, one empty tab per question."""
import warnings
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
warnings.filterwarnings("ignore")

orig = pd.read_pickle("work/orig.pkl")
final = pd.read_pickle("work/final.pkl")
COLS = ["Store", "Country", "Region", "Date", "Item", "Salesperson", "List Price", "Actual Price", "Discount %"]
HDR = PatternFill("solid", fgColor="DDEBF7")

qs = load_workbook("input/A_questions.xlsx", data_only=True)["Questions"]
questions = {}
for r in range(1, 60):
    n = qs.cell(r, 1).value
    if isinstance(n, (int, float)):
        text = [qs.cell(r, 11).value]
        for rr in range(r + 1, r + 3):
            nxt = qs.cell(rr, 11).value
            if qs.cell(rr, 1).value is None and isinstance(nxt, str) and nxt.startswith(("Additionally", "Rename")):
                text.append(nxt)
        questions[int(n)] = " ".join(t.strip() for t in text if t)

wb = Workbook()
ws = wb.active
ws.title = "Questions"
ws["A1"] = "Home appliances retailer: pivot table assignment"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = ("Data_Original = the data as supplied (regions APAC, EMEA, LATAM, NA). It feeds Q1-Q7. "
            "Data = the same rows after Q8 (NA + LATAM renamed to Americas) with the 12 omitted APAC/EMEA "
            "transactions from Q9 added at the bottom (rows 2341-2352). It feeds Q8-Q11.")
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[2].height = 48
ws["A4"], ws["B4"] = "Q", "Question"
for c in ("A4", "B4"):
    ws[c].font = Font(bold=True); ws[c].fill = HDR
for i, n in enumerate(sorted(questions), start=5):
    ws.cell(i, 1, n)
    ws.cell(i, 2, questions[n]).alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 5
ws.column_dimensions["B"].width = 120

def data_sheet(name, df):
    s = wb.create_sheet(name)
    s.append(COLS)
    for c in s[1]:
        c.font = Font(bold=True); c.fill = HDR
    for row in df[COLS].itertuples(index=False):
        s.append([row[0], row[1], row[2], row[3].to_pydatetime(), row[4], row[5], int(row[6]), int(row[7]), float(row[8])])
    for r in range(2, len(df) + 2):
        s.cell(r, 4).number_format = "yyyy-mm-dd"
        s.cell(r, 9).number_format = "0.00%"
    for col, w in zip("ABCDEFGHI", (17, 15, 10, 12, 17, 22, 10, 12, 11)):
        s.column_dimensions[col].width = w
    s.freeze_panes = "A2"
    return s

data_sheet("Data_Original", orig)
d = data_sheet("Data", final)
added = PatternFill("solid", fgColor="FFF2CC")
for r in range(len(orig) + 2, len(final) + 2):
    for c in range(1, 10):
        d.cell(r, c).fill = added
d.cell(len(final) + 3, 1, "Rows 2341-2352 (shaded) are the 12 transactions added for Q9.").font = Font(italic=True)

for n in range(1, 12):
    s = wb.create_sheet(f"Q{n}")
    s["A1"] = f"Q{n}. {questions[n]}"
    s["A1"].font = Font(bold=True)
    s["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    s.merge_cells("A1:H1")
    s.row_dimensions[1].height = 45
    s.column_dimensions["A"].width = 22
    s.column_dimensions["B"].width = 18
    s.column_dimensions["C"].width = 18
    s.column_dimensions["D"].width = 16
    s.column_dimensions["E"].width = 14
    s.column_dimensions["F"].width = 14
NOTES = {
    1: ["How it is built", "Rows: Store (each store is a city)", "Values: Sum of Actual Price ($)",
        "Total sales: $678,016 across 50 stores (2,339 transactions)"],
    2: ["How it is built", "Rows: Item", "Values: Count of Item (one row in the data = one item sold)",
        "Total: 2,339 items sold"],
    3: ["How it is built", "Rows: Item   Columns: Region", "Values: Sum of Actual Price ($)"],
    4: ["How it is built", "Same pivot as Q3.", "Item filter: only Microwave, Oven, Refrigerator ticked",
        "Region filter: NA unticked", "Revenue shown: $211,877"],
    5: ["Answer", "Report filter (external filter): Country = Columbia", "Rows: Salesperson",
        "Values: Average of Discount %, sorted smallest to largest",
        "Best sales agent: Gary Percival (average discount 6.83%).",
        "He gives away the least discount, so his sales come closest to list price.",
        "Note: he made only 2 sales in Columbia; David Finnie is next (8.16% over 6 sales)."],
    6: ["How it is built", "Rows: Region > Country > Store (tabular form, labels in separate columns)",
        "Values: Sum of Actual Price ($)", "Subtotals for Country and Store switched off"],
    7: ["Answer", "Keep Region, Country and Store in Rows and collapse the Region field:",
        "Excel: right-click any Region label > Expand/Collapse > Collapse Entire Field",
        "(or click the - button next to each Region).",
        "Google Sheets: click the - next to each Region value.",
        "Nothing is removed; Expand Entire Field brings the countries and stores back."],
    8: ["How it was done", "In the Data sheet, Region column (C): Find & Replace with 'Match entire cell contents'",
        "  NA -> Americas, then LATAM -> Americas (all 745 rows).", "Then Refresh the pivot (PivotTable Analyze > Refresh).",
        "Data_Original keeps the old values so Q1-Q7 still show NA/LATAM.",
        "This pivot reads Data!A1:I2340, i.e. before the Q9 rows were added."],
    9: ["How it was done", "The 12 omitted rows were pasted under the last row of Data (rows 2341-2352).",
        "Discount % was not copied; it is (List - Actual) / List, matching the sheet.",
        "PivotTable Analyze > Change Data Source: Data!$A$1:$I$2352, then Refresh.",
        "APAC rises by $764 (5 rows) to $167,794; EMEA by $637 (7 rows) to $306,351.",
        "Total is now $679,417 from 2,351 transactions."],
    10: ["How it is built", "Report filter: Country = USA   Rows: Store   Values: Sum of Actual Price",
         "Pie chart created from this pivot (8 USA stores, $90,192).",
         "To make it a PivotChart in Excel: click inside the pivot > PivotTable Analyze > PivotChart > Pie."],
    11: ["Steps (a slicer cannot be generated outside Excel/Sheets, so add it on your own copy)",
         "Excel: first turn the Q10 chart into a PivotChart (Q10 tab note), then click inside the Q10 pivot",
         "> Insert > Slicer > tick Item > OK. Clicking an item filters the pivot and its PivotChart.",
         "Google Sheets: click the pivot > Data > Add a slicer > choose column Item."],
}
for n, lines in NOTES.items():
    s = wb[f"Q{n}"]
    for i, line in enumerate(lines):
        c = s.cell(3 + i, 8, line)
        if i == 0: c.font = Font(bold=True)
        if n == 5 and line.startswith("Best"): c.font = Font(bold=True, color="C00000")
wb.save("work/base.xlsx")
print("saved", len(orig), len(final), questions[4][:80], "|", questions[8][:60])
