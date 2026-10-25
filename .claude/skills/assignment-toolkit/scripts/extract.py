#!/usr/bin/env python3
"""Dump the text of a question or answer sheet so every line can be read.

Usage: python3 -I extract.py <file> [out_dir]

Handles .pdf, .docx, .xlsx/.csv, .txt/.md and images. For PDFs it also
rasterises every page to PNG in out_dir so scanned pages, tables, figures
and handwriting can be checked visually (no OCR engine is installed, so
image-only pages must be read from those PNGs).
"""
import csv
import os
import subprocess
import sys


def pdf(path, out_dir):
    import pdfplumber

    empty = []
    with pdfplumber.open(path) as doc:
        for i, page in enumerate(doc.pages, 1):
            text = page.extract_text() or ""
            print(f"\n===== PAGE {i} =====")
            print(text if text.strip() else "[no text layer: read the page image]")
            if not text.strip():
                empty.append(i)
            for t, table in enumerate(page.extract_tables(), 1):
                print(f"--- table {t} on page {i} ---")
                for row in table:
                    print(" | ".join("" if c is None else c for c in row))
    os.makedirs(out_dir, exist_ok=True)
    prefix = os.path.join(out_dir, os.path.splitext(os.path.basename(path))[0])
    subprocess.run(["pdftoppm", "-r", "110", "-png", path, prefix], check=True)
    print(f"\n[page images written to {out_dir}]")
    if empty:
        print(f"[pages with no text layer: {empty}]")


def docx(path):
    import docx as d

    doc = d.Document(path)
    for p in doc.paragraphs:
        if p.text.strip():
            style = p.style.name if p.style is not None else ""
            tag = f"[{style}] " if style.startswith(("Heading", "List")) else ""
            print(tag + p.text)
    for t, table in enumerate(doc.tables, 1):
        print(f"--- table {t} ---")
        for row in table.rows:
            print(" | ".join(c.text.strip() for c in row.cells))
    if doc.inline_shapes:
        print(f"[{len(doc.inline_shapes)} embedded image(s): convert to PDF and view pages if they may hold questions]")


def sheet(path):
    if path.endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                print(" | ".join(row))
        return
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        print(f"===== SHEET {ws.title} =====")
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                print(" | ".join("" if c is None else str(c) for c in row))


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(path) or ".", "_pages")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        pdf(path, out_dir)
    elif ext == ".docx":
        docx(path)
    elif ext in (".xlsx", ".xlsm", ".csv"):
        sheet(path)
    elif ext in (".txt", ".md"):
        with open(path, encoding="utf-8", errors="replace") as f:
            print(f.read())
    elif ext in (".png", ".jpg", ".jpeg", ".webp", ".gif"):
        print("[image file: open it with the Read tool and transcribe it visually]")
    elif ext in (".doc", ".odt", ".rtf", ".pptx"):
        os.makedirs(out_dir, exist_ok=True)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, path], check=True)
        pdf(os.path.join(out_dir, os.path.splitext(os.path.basename(path))[0] + ".pdf"), out_dir)
    else:
        sys.exit(f"unsupported file type: {ext}")


if __name__ == "__main__":
    main()
